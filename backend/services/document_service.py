from __future__ import annotations

import hashlib
import json
import shutil
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from backend.models.task import TaskDocument

MAX_DOCUMENTS = 5
MAX_DOCUMENT_SIZE = 10 * 1024 * 1024
MAX_TOTAL_SIZE = 25 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 50_000
SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}


class DocumentValidationError(ValueError):
    pass


@dataclass(frozen=True)
class DocumentUpload:
    filename: str
    content_type: str
    content: bytes


@dataclass(frozen=True)
class StoredDocument:
    metadata: TaskDocument
    extracted_text: str

    def artifact_content(self) -> str:
        return json.dumps(
            {
                **self.metadata.model_dump(),
                "extracted_text": self.extracted_text,
            },
            ensure_ascii=False,
        )


class DocumentService:
    def __init__(self, storage_root: str | Path) -> None:
        self.storage_root = Path(storage_root).resolve()

    def save_documents(
        self, task_id: str, uploads: list[DocumentUpload]
    ) -> list[StoredDocument]:
        prepared = self._prepare(uploads)
        task_directory = self.storage_root / task_id
        try:
            task_directory.mkdir(parents=True, exist_ok=False)
            stored: list[StoredDocument] = []
            for index, (upload, filename, extracted_text) in enumerate(
                prepared, start=1
            ):
                destination = task_directory / f"{index:02d}-{filename}"
                destination.write_bytes(upload.content)
                metadata = TaskDocument(
                    filename=filename,
                    content_type=upload.content_type
                    or "application/octet-stream",
                    size=len(upload.content),
                    sha256=hashlib.sha256(upload.content).hexdigest(),
                    storage_path=str(destination),
                )
                stored.append(
                    StoredDocument(
                        metadata=metadata,
                        extracted_text=extracted_text,
                    )
                )
            return stored
        except Exception:
            shutil.rmtree(task_directory, ignore_errors=True)
            raise

    def validate_documents(self, uploads: list[DocumentUpload]) -> None:
        self._prepare(uploads)

    def _prepare(
        self, uploads: list[DocumentUpload]
    ) -> list[tuple[DocumentUpload, str, str]]:
        if not uploads:
            return []
        if len(uploads) > MAX_DOCUMENTS:
            raise DocumentValidationError(
                f"最多只能上传 {MAX_DOCUMENTS} 个文档"
            )
        if sum(len(item.content) for item in uploads) > MAX_TOTAL_SIZE:
            raise DocumentValidationError("文档总大小不能超过 25 MB")

        prepared: list[tuple[DocumentUpload, str, str]] = []
        filenames: set[str] = set()
        for upload in uploads:
            filename = upload.filename.replace("\\", "/").split("/")[-1]
            if not filename or filename in {".", ".."}:
                raise DocumentValidationError("文档文件名无效")
            if filename in filenames:
                raise DocumentValidationError(f"文档文件名重复：{filename}")
            filenames.add(filename)
            extension = Path(filename).suffix.lower()
            if extension not in SUPPORTED_EXTENSIONS:
                raise DocumentValidationError(
                    f"不支持的文档格式：{extension or '无扩展名'}"
                )
            if not upload.content:
                raise DocumentValidationError(f"文档为空：{filename}")
            if len(upload.content) > MAX_DOCUMENT_SIZE:
                raise DocumentValidationError(
                    f"文档超过 10 MB：{filename}"
                )
            extracted = self._extract_text(extension, upload.content)
            prepared.append(
                (
                    upload,
                    filename,
                    extracted[:MAX_EXTRACTED_CHARACTERS],
                )
            )
        return prepared

    @staticmethod
    def _extract_text(extension: str, content: bytes) -> str:
        try:
            if extension in {".txt", ".md", ".csv", ".json"}:
                return content.decode("utf-8-sig", errors="replace")
            if extension == ".pdf":
                reader = PdfReader(BytesIO(content))
                return "\n\n".join(
                    page.extract_text() or "" for page in reader.pages
                )
            if extension == ".docx":
                document = Document(BytesIO(content))
                return "\n".join(
                    paragraph.text for paragraph in document.paragraphs
                )
        except Exception as exc:
            raise DocumentValidationError(
                f"无法读取文档内容：{exc}"
            ) from exc
        return ""
